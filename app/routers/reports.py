from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.database import get_db
from app.dependencies import require_role
from app.models import UserRole
from app.schemas import DailyTechnicianReportOut, MonthlyReportOut
from app.services.report_service import get_daily_technician_report, get_monthly_report

router = APIRouter(prefix="/reports", tags=["reports"])


@router.get("/monthly", response_model=MonthlyReportOut)
def monthly_report(
    year: int = None,
    month: int = None,
    db: Session = Depends(get_db),
    current_user=Depends(require_role(UserRole.ADMIN)),
):
    now = datetime.utcnow()
    year = year or now.year
    month = month or now.month
    return get_monthly_report(year, month, db)


@router.get("/daily-technicians", response_model=list[DailyTechnicianReportOut])
def daily_technician_report(
    report_date: date = None,
    db: Session = Depends(get_db),
    current_user=Depends(require_role(UserRole.ADMIN)),
):
    report_date = report_date or date.today()
    return get_daily_technician_report(report_date, db)


@router.get("/daily-technicians/export")
def export_daily_technician_report(
    report_date: date = None,
    db: Session = Depends(get_db),
    current_user=Depends(require_role(UserRole.ADMIN)),
):
    report_date = report_date or date.today()
    entries = get_daily_technician_report(report_date, db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Technician Report"

    headers = ["Technician", "Regional Center", "Jobs Closed", "Total Hours", "Total KM", "Job Numbers"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_num, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for entry in entries:
        ws.append([
            entry["technician_name"],
            entry["regional_center"] or "",
            entry["jobs_closed"],
            entry["total_hours"],
            entry["total_km"],
            ", ".join(entry["job_numbers"]),
        ])

    if not entries:
        ws.append(["No jobs were closed on this date.", "", "", "", "", ""])

    column_widths = [24, 18, 12, 12, 10, 40]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"daily_technician_report_{report_date.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )